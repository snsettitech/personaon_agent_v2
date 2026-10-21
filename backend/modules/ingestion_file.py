"""
File-based text extraction for ingestion.

Extracted from ingestion.py -- contains PDF, DOCX, and Excel extraction logic.
"""

import os
import re
from typing import Any

from PyPDF2 import PdfReader
import docx
import openpyxl

from modules.doc_sectioning import extract_section_blocks
from modules.ingestion_config import _to_bool


# Imported lazily to avoid circular imports; chunk_text is defined in ingestion_chunking.
def _get_chunk_text():
    from modules.ingestion_chunking import chunk_text
    return chunk_text


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from PDF. Tries Gemini native PDF reading first (handles
    scanned docs, images, tables), falls back to PyPDF2.
    """
    import asyncio
    import logging as _logging
    _logger = _logging.getLogger(__name__)

    try:
        from modules.gemini_extraction import gemini_extract_pdf
        try:
            loop = asyncio.get_event_loop()
            if loop.is_running():
                import concurrent.futures
                with concurrent.futures.ThreadPoolExecutor() as pool:
                    future = pool.submit(asyncio.run, gemini_extract_pdf(file_path))
                    gemini_text = future.result(timeout=120)
            else:
                gemini_text = loop.run_until_complete(gemini_extract_pdf(file_path))
        except RuntimeError:
            gemini_text = asyncio.run(gemini_extract_pdf(file_path))

        if gemini_text:
            _logger.info("[PDF] Gemini extraction succeeded: %d chars", len(gemini_text))
            return gemini_text
    except Exception as e:
        _logger.warning("[PDF] Gemini extraction failed, falling back to PyPDF2: %s", e)

    text, _chunk_entries = extract_pdf_text_and_chunk_entries(
        file_path,
        doc_name=os.path.basename(file_path) or "document",
    )
    return text


def _extract_pdf_pages(file_path: str) -> list[dict[str, Any]]:
    """
    Extract per-page text from a PDF while preserving page numbers.
    """
    reader = PdfReader(file_path)
    pages: list[dict[str, Any]] = []
    for page_number, page in enumerate(reader.pages, 1):
        # Some PDFs return None for image-only pages; keep extraction best-effort
        # instead of crashing the whole ingest on a single page.
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(
                {
                    "page_number": page_number,
                    "text": page_text,
                }
            )
    return pages


def _safe_doc_name(value: str) -> str:
    name = re.sub(r"\s+", " ", str(value or "").strip())
    if not name:
        return "document"
    return name.replace("\\", "/")


def _build_pdf_chunk_entries(
    pages: list[dict[str, Any]],
    *,
    doc_name: str,
    chunk_size: int = 1000,
    overlap: int = 200,
) -> list[dict[str, Any]]:
    chunk_text_fn = _get_chunk_text()
    entries: list[dict[str, Any]] = []
    safe_doc_name = _safe_doc_name(doc_name)

    for page in pages:
        page_number = page.get("page_number")
        page_text = str(page.get("text") or "").strip()
        if not page_text:
            continue

        page_tag = f"page_{page_number}"
        fallback_section_title = page_tag
        fallback_section_path = f"{safe_doc_name}/{page_tag}"

        section_blocks = extract_section_blocks(
            page_text,
            base_metadata={"page_number": page_number},
        )
        if not section_blocks:
            section_blocks = [
                {
                    "text": page_text,
                    "section_title": fallback_section_title,
                    "section_path": fallback_section_path,
                    "chunk_type": "page",
                    "block_type": "answer_text",
                    "is_answer_text": True,
                    "page_number": page_number,
                }
            ]

        for block in section_blocks:
            block_text = str(block.get("text") or "").strip()
            if not block_text:
                continue

            raw_title = str(block.get("section_title") or "").strip()
            raw_path = str(block.get("section_path") or "").strip()

            # For PDFs, avoid generic section labels and always keep page fallback.
            section_title = (
                fallback_section_title
                if not raw_title or raw_title.lower() in {"general", "unknown"}
                else raw_title
            )
            section_path = raw_path
            if not section_path or section_path.lower() in {"general", "unknown"}:
                section_path = fallback_section_path
            elif not section_path.lower().startswith(f"{safe_doc_name.lower()}/"):
                section_path = f"{safe_doc_name}/{section_path}"

            for chunk in chunk_text_fn(block_text, chunk_size=chunk_size, overlap=overlap):
                if not chunk or not chunk.strip():
                    continue
                entries.append(
                    {
                        "text": chunk,
                        "section_title": section_title,
                        "section_path": section_path,
                        "chunk_type": str(block.get("chunk_type") or "section"),
                        "block_type": str(block.get("block_type") or "answer_text"),
                        "is_answer_text": _to_bool(block.get("is_answer_text"), default=True),
                        "page_number": page_number,
                    }
                )

    return entries


def extract_pdf_text_and_chunk_entries(
    file_path: str,
    *,
    doc_name: str,
    chunk_size: int = 1000,
    overlap: int = 200,
) -> tuple[str, list[dict[str, Any]]]:
    pages = _extract_pdf_pages(file_path)
    text = "\n".join(str(page.get("text") or "") for page in pages if page.get("text"))
    chunk_entries = _build_pdf_chunk_entries(
        pages,
        doc_name=doc_name,
        chunk_size=chunk_size,
        overlap=overlap,
    )
    return text, chunk_entries


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a Word document."""
    try:
        doc = docx.Document(file_path)
        text = []
        for para in doc.paragraphs:
            if para.text:
                text.append(para.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text.append(cell.text)

        result = "\n".join(text)
        if not result.strip():
            raise ValueError("DOCX file appears to be empty or contains no extractable text")
        return result
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_excel(file_path: str) -> str:
    """Extract text from all sheets in an Excel file."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        text.append(f"--- Sheet: {sheet_name} ---")
        for row in sheet.iter_rows(values_only=True):
            # Filter out None values and convert to string
            row_text = [str(cell) for cell in row if cell is not None]
            if row_text:
                text.append(" | ".join(row_text))
    return "\n".join(text)
